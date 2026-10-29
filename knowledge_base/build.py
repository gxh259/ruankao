"""Index the local 2026-05 高项 course material without copying it to Git."""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import sqlite3
from concurrent.futures import ProcessPoolExecutor, as_completed
from pathlib import Path

SOURCE = Path(os.environ.get('RUANKAO_SOURCE',
    Path.home() / 'Desktop' / '软考【高项】2026年5月班备考资料'))
LOCAL = Path(__file__).resolve().parent / 'local'


def extract(path_str: str) -> tuple[str, list[str], str | None]:
    path = Path(path_str)
    try:
        if path.suffix.lower() == '.pdf':
            from pypdf import PdfReader
            reader = PdfReader(path, strict=False)
            pages = [(page.extract_text() or '').strip() for page in reader.pages]
        elif path.suffix.lower() == '.docx':
            from docx import Document
            doc = Document(path)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            table_rows = [
                '\t'.join(cell.text.replace('\n', ' ') for cell in row.cells)
                for table in doc.tables for row in table.rows
            ]
            pages = ['\n'.join(paragraphs + table_rows)]
        elif path.suffix.lower() == '.xlsx':
            from openpyxl import load_workbook
            book = load_workbook(path, read_only=True, data_only=True)
            pages = []
            for sheet in book:
                rows = ['\t'.join(str(v) if v is not None else '' for v in row)
                        for row in sheet.values]
                pages.append(sheet.title + '\n' + '\n'.join(rows))
            book.close()
        else:
            return path_str, [], 'unsupported format'
        return path_str, pages, None
    except Exception as exc:
        return path_str, [], f'{type(exc).__name__}: {exc}'


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument('--workers', type=int, default=4)
    args = parser.parse_args()
    LOCAL.mkdir(parents=True, exist_ok=True)
    db = sqlite3.connect(LOCAL / 'kb.sqlite')
    db.executescript('''
        PRAGMA journal_mode=WAL;
        DROP TABLE IF EXISTS pages;
        DROP TABLE IF EXISTS documents;
        CREATE TABLE documents (id TEXT PRIMARY KEY, path TEXT UNIQUE, category TEXT,
            extension TEXT, bytes INTEGER, page_count INTEGER, text_chars INTEGER,
            status TEXT, error TEXT);
        CREATE TABLE pages (document_id TEXT, page_number INTEGER, text TEXT,
            PRIMARY KEY(document_id, page_number));
        CREATE INDEX pages_doc ON pages(document_id);
    ''')
    paths = sorted(p for p in SOURCE.rglob('*') if p.is_file() and p.name != '.DS_Store')
    searchable = [p for p in paths if p.suffix.lower() in {'.pdf', '.docx', '.xlsx'}]
    for path in paths:
        rel = str(path.relative_to(SOURCE))
        doc_id = hashlib.sha256(rel.encode()).hexdigest()[:16]
        db.execute('INSERT INTO documents VALUES (?,?,?,?,?,?,?,?,?)',
                   (doc_id, rel, rel.split('/')[0], path.suffix.lower(),
                    path.stat().st_size, 0, 0, 'pending' if path in searchable else 'media_or_unsupported', None))
    db.commit()
    with ProcessPoolExecutor(max_workers=args.workers) as pool:
        jobs = {pool.submit(extract, str(p)): p for p in searchable}
        for i, job in enumerate(as_completed(jobs), 1):
            path_str, pages, error = job.result()
            path = Path(path_str)
            rel = str(path.relative_to(SOURCE))
            doc_id = hashlib.sha256(rel.encode()).hexdigest()[:16]
            db.executemany('INSERT INTO pages VALUES (?,?,?)',
                           ((doc_id, n, page) for n, page in enumerate(pages, 1) if page))
            chars = sum(map(len, pages))
            status = 'error' if error else ('text' if chars >= 100 else 'little_or_no_text')
            db.execute('UPDATE documents SET page_count=?, text_chars=?, status=?, error=? WHERE id=?',
                       (len(pages), chars, status, error, doc_id))
            if i % 25 == 0:
                db.commit()
                print(f'{i}/{len(searchable)} indexed', flush=True)
    db.commit()
    summary = dict(db.execute('SELECT status, count(*) FROM documents GROUP BY status'))
    print(json.dumps(summary, ensure_ascii=False))
    print('database:', LOCAL / 'kb.sqlite')


if __name__ == '__main__':
    main()
